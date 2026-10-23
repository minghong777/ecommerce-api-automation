from dataclasses import replace
from datetime import date, datetime, timezone
from io import BytesIO, StringIO
import csv
import json
from decimal import Decimal
import httpx
from openpyxl import load_workbook
import pytest
from src.api_client import APIError, fetch_products
from src.demo_data import generate_dataset, offline_products
from src.metrics import Filters, calculate, daily_sales, filter_dataset, top_products
from src.models import Customer, Dataset, Order, OrderItem, Product, ValidationError, cents
from src.reports import export_csv, export_excel, export_json, report, safe_cell
from src.transform import parse_products

DAY = datetime(2026, 7, 1, tzinfo=timezone.utc)
FILTER = Filters(date(2026, 7, 1), date(2026, 9, 30))

@pytest.fixture
def small():
    products = (Product(1, 'Demo One', 'demo', 99999), Product(2, '=HYPERLINK("bad")', 'demo', 100))
    orders = tuple(Order(f'O{i}', 'C1', DAY, s, 50, 10) for i, s in enumerate(('paid', 'refunded', 'cancelled', 'pending')))
    items = tuple(OrderItem(o.order_id, 1, 2, 100) for o in orders)
    data = Dataset(products, (Customer('C1', 'Demo Customer'),), orders, items)
    data.validate()
    return data


def test_deterministic_offline():
    one = generate_dataset()
    assert one == generate_dataset()
    assert one != generate_dataset(seed=42)
    one.validate()
    assert len(one.orders) == 240
    assert all(c.display_name.startswith('Demo Customer') for c in one.customers)
    assert {i.product_id for i in one.items} <= {p.product_id for p in offline_products()}

@pytest.mark.parametrize('raw,expected', [('1.005', 101), ('2.675', 268), ('0.1', 10), ('0', 0)])
def test_rounding(raw, expected):
    assert cents(raw) == expected

@pytest.mark.parametrize('raw', [-1, 'NaN', 'Infinity', True, None, '1000001'])
def test_invalid_money(raw):
    with pytest.raises(ValidationError): cents(raw)


def test_financial_policy(small):
    metrics = calculate(small)
    assert metrics == {'gmv':800, 'discounts':10, 'shipping_fees':50, 'refunds':240,
                       'cancelled_orders':1, 'cancelled_gmv':200, 'refunded_orders':1,
                       'paid_gmv':200, 'net_sales':240, 'paid_order_count':1,
                       'average_order_value':240, 'units_sold':2, 'selected_order_count':4}
    assert top_products(small)[0]['net_revenue'] == 190
    # Catalog price 999.99 has no influence on historical purchase price 1.00.
    assert calculate(replace(small, products=(replace(small.products[0], unit_price=1), small.products[1]))) == metrics


def test_filtering_and_empty(small):
    filtered = filter_dataset(small, replace(FILTER, statuses=('refunded',), search='o1'))
    assert [o.order_id for o in filtered.orders] == ['O1']
    assert len(filtered.items) == 1
    empty = filter_dataset(small, replace(FILTER, start=date(2026, 8, 1), statuses=()))
    assert not empty.items and not empty.orders
    assert calculate(empty)['average_order_value'] == 0
    assert sum(r['net_sales'] for r in daily_sales(empty, FILTER)) == 0
    with pytest.raises(ValueError): filter_dataset(small, replace(FILTER, end=date(2026, 1, 1)))


def test_allocation_reconciles():
    data = generate_dataset()
    metrics = calculate(data)
    assert sum(p['net_revenue'] for p in top_products(data)) == metrics['net_sales'] - metrics['shipping_fees']
    assert sum(p['units_sold'] for p in top_products(data)) == metrics['units_sold']
    assert sum(d['net_sales'] for d in daily_sales(data, FILTER)) == metrics['net_sales']

@pytest.mark.parametrize('change', ['duplicate_order', 'customer_fk', 'product_fk', 'quantity', 'price', 'status', 'date', 'discount', 'missing_items'])
def test_validation(small, change):
    if change == 'duplicate_order': data = replace(small, orders=small.orders + small.orders[:1])
    elif change == 'customer_fk': data = replace(small, customers=())
    elif change == 'product_fk': data = replace(small, products=small.products[1:])
    elif change == 'quantity': data = replace(small, items=(replace(small.items[0], quantity=0),) + small.items[1:])
    elif change == 'price': data = replace(small, items=(replace(small.items[0], unit_price_at_purchase=-1),) + small.items[1:])
    elif change == 'status': data = replace(small, orders=(replace(small.orders[0], status='fake'),) + small.orders[1:])
    elif change == 'date': data = replace(small, orders=(replace(small.orders[0], order_date=DAY.replace(tzinfo=None)),) + small.orders[1:])
    elif change == 'discount': data = replace(small, orders=(replace(small.orders[0], discount_amount=201),) + small.orders[1:])
    else: data = replace(small, items=small.items[1:])
    with pytest.raises(ValidationError): data.validate()


def test_exports(small):
    filtered = filter_dataset(small, replace(FILTER, statuses=('paid',)))
    payload = report(filtered, replace(FILTER, statuses=('paid',)), 'Offline fictional')
    parsed = json.loads(export_json(payload))
    assert parsed['summary']['net_sales'] == '2.40'
    assert len(parsed['orders']) == len(parsed['order_items']) == 1
    csv_rows = list(csv.DictReader(StringIO(export_csv(payload).decode('utf-8-sig'))))
    assert csv_rows[0]['net_sales'] == '2.40'
    workbook = load_workbook(BytesIO(export_excel(payload)))
    assert workbook.sheetnames == ['Summary', 'Orders', 'OrderItems', 'TopProducts']
    assert workbook['Orders'].max_row == 2
    assert workbook['OrderItems']['D2'].value == '1.00'
    assert Decimal(csv_rows[0]['net_sales']) == Decimal(parsed['summary']['net_sales'])

@pytest.mark.parametrize('text', ['=SUM(A1)', '+cmd', '-cmd', '@SUM(A1)', '  =1', '\tunsafe', '\runsafe'])
def test_safe_cells(text):
    assert safe_cell(text).startswith("'")


def test_excel_untrusted_product(small):
    data = replace(small, items=tuple(replace(i, product_id=2) for i in small.items))
    workbook = load_workbook(BytesIO(export_excel(report(data, FILTER, '=untrusted'))))
    assert workbook['TopProducts']['B2'].value.startswith("'")
    assert not any(c.data_type == 'f' for sheet in workbook for row in sheet for c in row)


def test_empty_exports(small):
    empty = filter_dataset(small, replace(FILTER, statuses=()))
    payload = report(empty, replace(FILTER, statuses=()), 'offline')
    assert len(export_csv(payload).decode('utf-8-sig').splitlines()) == 1
    wb = load_workbook(BytesIO(export_excel(payload)))
    assert wb['Orders'].max_row == 1
    assert json.loads(export_json(payload))['summary']['paid_order_count'] == 0


def test_api_pagination():
    calls = []
    def handler(request):
        calls.append(request)
        skip = int(request.url.params['skip'])
        return httpx.Response(200, json={'products':[{'id':skip+1,'title':'Demo','category':'demo','price':1.005}], 'total':2,'skip':skip})
    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        result = fetch_products(client, page_size=1)
    assert len(result.products) == result.total == result.requests == 2
    assert result.products[0].unit_price == 101
    assert calls[1].url.params['skip'] == '1'
    assert result.fetched_at.tzinfo is not None

@pytest.mark.parametrize('failure', ['timeout', '503', '429'])
def test_api_retry_bounded(failure):
    calls, delays = [], []
    def handler(request):
        calls.append(request)
        if failure == 'timeout': raise httpx.ReadTimeout('timeout')
        return httpx.Response(int(failure))
    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        with pytest.raises(APIError): fetch_products(client, sleep=delays.append)
    assert len(calls) == 3 and delays == [0.25, 0.5]


def test_api_transient_recovery():
    calls = []
    def handler(request):
        calls.append(request)
        return httpx.Response(503) if len(calls) == 1 else httpx.Response(200, json={'products':[{'id':1,'title':'Demo','category':'demo','price':1}], 'total':1,'skip':0})
    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        assert fetch_products(client, sleep=lambda _: None).requests == 2

@pytest.mark.parametrize('payload', [{}, {'products':[{'id':True}]}, {'products':[{'id':1,'title':'Demo','category':'x','price':-1}]}])
def test_bad_payload(payload):
    with pytest.raises(ValidationError): parse_products(payload)


def test_bad_http_and_json():
    for response in [httpx.Response(404), httpx.Response(200, text='not JSON')]:
        with httpx.Client(transport=httpx.MockTransport(lambda r: response)) as client:
            with pytest.raises(APIError): fetch_products(client)


def test_api_repeated_page():
    with httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(200, json={
            'products':[{'id':1,'title':'Demo','category':'demo','price':1}], 'total':2,'skip':int(r.url.params['skip'])}))) as client:
        with pytest.raises(APIError): fetch_products(client, page_size=1)


def test_streamlit_offline_pages():
    from streamlit.testing.v1 import AppTest
    from src.config import ROOT
    app = AppTest.from_file(ROOT / 'app.py').run(timeout=30)
    assert not app.exception
    assert len(app.metric) == 4
    for page in ['Orders Explorer', 'Sales Analytics', 'Report Export', 'Integrations (Coming Soon)', 'API Connection / Data Source']:
        app.sidebar.radio[0].set_value(page).run(timeout=30)
        assert not app.exception, page
    app.sidebar.multiselect[0].set_value([]).run()
    assert not app.exception

@pytest.mark.parametrize('status', ['paid', 'shipped', 'completed'])
def test_recognized_statuses(small, status):
    data = replace(small, orders=(replace(small.orders[0], status=status),), items=small.items[:1])
    assert calculate(data)['net_sales'] == 240


def test_discount_remainder_tie(small):
    data = replace(small, orders=(replace(small.orders[0], discount_amount=1),),
                   items=(OrderItem('O0', 1, 1, 100), OrderItem('O0', 2, 1, 100)))
    data.validate()
    rows = {r['product_id']: r for r in top_products(data)}
    assert rows[1]['net_revenue'] == 99
    assert rows[2]['net_revenue'] == 100


def test_api_size_guard():
    from src.config import MAX_RESPONSE_BYTES
    with httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(200, content=b'x' * (MAX_RESPONSE_BYTES + 1)))) as client:
        with pytest.raises(APIError, match='size'): fetch_products(client)


def test_api_cap():
    def handler(request):
        skip, limit = int(request.url.params['skip']), int(request.url.params['limit'])
        return httpx.Response(200, json={'products':[{'id':i+1,'title':'Demo','category':'x','price':1} for i in range(skip, skip+limit)], 'total':1000,'skip':skip})
    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        result = fetch_products(client, page_size=2, max_products=3)
    assert len(result.products) == 3 and result.requests == 2


def test_utc_date_boundary(small):
    from datetime import timedelta
    # Local July 2 at 01:00 is still July 1 in UTC.
    order = replace(small.orders[0], order_date=datetime(2026, 7, 2, 1, tzinfo=timezone(timedelta(hours=8))))
    data = replace(small, orders=(order,), items=small.items[:1])
    filtered = filter_dataset(data, replace(FILTER, end=date(2026, 7, 1)))
    assert len(filtered.orders) == 1
    assert daily_sales(filtered, replace(FILTER, end=date(2026, 7, 1)))[0]['net_sales'] == 240


def test_aov_half_up(small):
    orders = (replace(small.orders[0], shipping_fee=0, discount_amount=0),
              replace(small.orders[1], status='paid', shipping_fee=0, discount_amount=0))
    data = replace(small, orders=orders, items=(OrderItem('O0', 1, 1, 100), OrderItem('O1', 1, 1, 101)))
    assert calculate(data)['average_order_value'] == 101
