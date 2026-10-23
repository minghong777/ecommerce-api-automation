"""In-memory reporting; money is exported as exact decimal strings."""
import csv
from dataclasses import asdict
from datetime import datetime, timezone
from io import BytesIO, StringIO
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from .config import MAX_EXPORT_BYTES
from .metrics import Filters, calculate, order_rows, top_products
from .models import Dataset

MONEY_FIELDS = {'unit_price', 'unit_price_at_purchase', 'shipping_fee', 'discount_amount',
                'gmv', 'refunds', 'net_sales', 'net_revenue'}
ORDER_COLUMNS = ['order_id', 'customer_id', 'order_date', 'status', 'shipping_fee', 'discount_amount', 'gmv', 'units', 'refunds', 'net_sales']
ITEM_COLUMNS = ['order_id', 'product_id', 'quantity', 'unit_price_at_purchase']
TOP_COLUMNS = ['product_id', 'title', 'category', 'units_sold', 'net_revenue']


def decimal_money(value: int) -> str:
    return f'{value // 100}.{value % 100:02d}'


def safe_cell(value: object) -> object:
    if isinstance(value, str):
        value = ''.join(c for c in value if ord(c) >= 32 or c in '\t\n\r')
        if value.lstrip().startswith(('=', '+', '-', '@')) or value.startswith(('\t', '\r', '\n')):
            return "'" + value
    return value


def report(data: Dataset, filters: Filters, provenance: str) -> dict:
    metrics = calculate(data)
    monetary = {'gmv', 'discounts', 'shipping_fees', 'refunds', 'cancelled_gmv', 'paid_gmv', 'net_sales', 'average_order_value'}
    def converted(rows: list[dict]) -> list[dict]:
        return [{k: decimal_money(v) if k in MONEY_FIELDS else v for k, v in row.items()} for row in rows]
    return {
        'metadata': {'disclosure': 'DEMO / SYNTHETIC DATA — independent portfolio project',
                     'currency': 'USD', 'money_format': 'Decimal strings in USD', 'source': provenance,
                     'generated_at': datetime.now(timezone.utc).isoformat(),
                     'filters': {**asdict(filters), 'start': filters.start.isoformat(), 'end': filters.end.isoformat()},
                     'revenue_rule': 'Paid/shipped/completed merchandise less discounts plus shipping; other statuses zero.'},
        'summary': {k: decimal_money(v) if k in monetary else v for k, v in metrics.items()},
        'orders': converted(order_rows(data)), 'order_items': converted([asdict(i) for i in data.items]),
        'top_products': converted(top_products(data)),
    }


def bounded(data: bytes) -> bytes:
    if len(data) > MAX_EXPORT_BYTES:
        raise ValueError('Report exceeds the 10 MB download limit.')
    return data


def export_json(payload: dict) -> bytes:
    return bounded(json.dumps(payload, indent=2, ensure_ascii=False).encode())


def export_csv(payload: dict) -> bytes:
    """Flat filtered orders; JSON and Excel include summary and line items."""
    text = StringIO(newline='')
    writer = csv.DictWriter(text, fieldnames=ORDER_COLUMNS)
    writer.writeheader()
    writer.writerows({k: safe_cell(v) for k, v in row.items()} for row in payload['orders'])
    return bounded(text.getvalue().encode('utf-8-sig'))


def export_excel(payload: dict) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Summary'
    summary.append(['DEMO / SYNTHETIC DATA', 'USD; monetary values are exact decimal text'])
    for key, value in payload['metadata'].items():
        summary.append([key, safe_cell(json.dumps(value) if isinstance(value, dict) else value)])
    for key, value in payload['summary'].items():
        summary.append([key, value])
    for name, key, columns in [('Orders', 'orders', ORDER_COLUMNS), ('OrderItems', 'order_items', ITEM_COLUMNS),
                                ('TopProducts', 'top_products', TOP_COLUMNS)]:
        sheet = workbook.create_sheet(name)
        sheet.append(columns)
        for row in payload[key]:
            sheet.append([safe_cell(row[c]) for c in columns])
    for sheet in workbook:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='087F8C')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(65, max(16, max(len(str(c.value or '')) for c in column) + 2))
    output = BytesIO()
    workbook.save(output)
    return bounded(output.getvalue())
